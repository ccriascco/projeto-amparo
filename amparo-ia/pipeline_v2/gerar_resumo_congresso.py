"""Gera o resumo para o congresso no modelo da UNIRP (título, autores, instituição, parágrafo único, palavras-chave)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

SAIDA = Path(__file__).resolve().parent / 'Resumo_Congresso_Classificacao_Risco_IA.docx'

TITULO = ('CLASSIFICAÇÃO DO RISCO DE VIOLÊNCIA CONTRA A MULHER COM INTELIGÊNCIA ARTIFICIAL: '
          'UM ESTUDO COM FLORESTA ALEATÓRIA')

AUTORES = [
    '[NOME COMPLETO DO AUTOR]',
    '[NOME COMPLETO DO COAUTOR]',
    '[NOME COMPLETO DO ORIENTADOR]',
]

INSTITUICAO = '[NOME COMPLETO DA INSTITUIÇÃO] (UNIRP)'

CORPO = (
    'A violência contra a mulher exige respostas rápidas e uma avaliação confiável do risco a que a vítima está '
    'exposta. Este trabalho tem como objetivo desenvolver e avaliar um modelo de inteligência artificial capaz de '
    'classificar o nível de risco de uma usuária do aplicativo Amparo, em baixo, médio ou alto, a partir do seu '
    'histórico de ocorrências e do uso do botão de emergência. A metodologia consistiu em definir regras de negócio '
    'que associam tipos de ocorrência, recorrência, recência e acionamentos válidos do botão a cada nível de risco, '
    'e em gerar 9.000 casos sintéticos rotulados por essas regras, distribuídos de forma equilibrada entre os três '
    'níveis. Os dados foram divididos em treino, validação e teste, sem exemplos repetidos entre os conjuntos. Para '
    'simular erros humanos de rotulagem, 10% dos rótulos de treino e validação foram alterados de forma aleatória, '
    'enquanto o teste permaneceu íntegro. O modelo escolhido foi a floresta aleatória, com 100 árvores, comparada a '
    'uma árvore de decisão única. Os resultados mostram que a floresta aleatória atingiu acurácia de 99,9% no teste, '
    'com F1-score de 0,999 e nenhum caso de risco alto classificado como baixo, enquanto a árvore de decisão cometeu '
    'um erro desse tipo. Ao elevar para 30% a fração de rótulos incorretos no treino, a acurácia da floresta aleatória '
    'se manteve em 99,8%, enquanto a da árvore caiu para 75%, o que evidencia maior robustez do primeiro. A análise '
    'também indicou que o acionamento do botão de emergência eleva o risco conforme a regra definida, sem ser '
    'supervalorizado pelo modelo. Conclui-se que a floresta aleatória é uma abordagem promissora para apoiar a '
    'classificação de risco no aplicativo. Contudo, por se tratar de dados sintéticos, os resultados indicam a '
    'capacidade de reproduzir as regras definidas e não garantem desempenho com casos reais. A validação com dados '
    'reais rotulados por profissionais especializados é a próxima etapa necessária antes de sua utilização.'
)

PALAVRAS_CHAVE = [
    'Inteligência Artificial',
    'Classificação de Risco',
    'Violência Contra a Mulher',
    'Floresta Aleatória',
    'Aprendizado de Máquina',
]


def main():
    tamanho = len(CORPO)
    if tamanho > 3000:
        raise SystemExit(f'Corpo com {tamanho} caracteres: acima do limite de 3.000')

    doc = Document()
    estilo = doc.styles['Normal']
    estilo.font.name = 'Times New Roman'
    estilo.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(TITULO)
    r.bold = True

    for autor in AUTORES:
        p = doc.add_paragraph(autor)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph(INSTITUICAO)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(CORPO).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p = doc.add_paragraph()
    p.add_run('Palavras-chave: ').bold = True
    p.add_run(', '.join(PALAVRAS_CHAVE) + '.')

    doc.save(SAIDA)
    print(f'Corpo com {tamanho} caracteres (limite 3.000).')
    print('Documento gerado em', SAIDA)


if __name__ == '__main__':
    main()
